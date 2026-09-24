"""Text extraction — the "text brain" half of §11.

PDF, TXT, Markdown, DOCX and images.  Extraction is per page so every chunk
keeps a page reference, which is what lets an answer say "page 47" and what
lets the visual brain open exactly that page.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.errors import UnsupportedMedia
from app.logging_setup import get_logger

log = get_logger(__name__)

SUPPORTED_MIME = {
    "application/pdf": "pdf",
    "text/plain": "text",
    "text/markdown": "text",
    "text/x-markdown": "text",
    "text/csv": "text",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "image/jpeg": "image",
    "image/png": "image",
    "image/webp": "image",
    "image/bmp": "image",
    "image/tiff": "image",
}
SUPPORTED_EXT = {
    ".pdf": "pdf", ".txt": "text", ".md": "text", ".markdown": "text", ".csv": "text",
    ".log": "text", ".json": "text", ".docx": "docx",
    ".jpg": "image", ".jpeg": "image", ".png": "image", ".webp": "image",
    ".bmp": "image", ".tif": "image", ".tiff": "image",
}


@dataclass
class Page:
    number: int
    text: str = ""
    char_count: int = 0
    needs_ocr: bool = False


@dataclass
class Extraction:
    kind: str
    pages: List[Page] = field(default_factory=list)
    page_count: int = 0
    total_chars: int = 0
    engine: str = "none"
    warnings: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "kind": self.kind, "page_count": self.page_count,
            "total_chars": self.total_chars, "engine": self.engine,
            "warnings": self.warnings,
            "pages": [{"number": p.number, "chars": p.char_count,
                       "needs_ocr": p.needs_ocr} for p in self.pages],
        }


def classify(path: Path, mime_type: Optional[str]) -> str:
    kind = SUPPORTED_MIME.get((mime_type or "").split(";")[0].strip().lower())
    if kind:
        return kind
    kind = SUPPORTED_EXT.get(path.suffix.lower())
    if kind:
        return kind
    raise UnsupportedMedia(
        f"'{path.suffix or mime_type or 'unknown'}' is not a supported document type",
        supported=sorted(set(SUPPORTED_EXT)),
    )


def extract(path: Path, mime_type: Optional[str] = None) -> Extraction:
    kind = classify(path, mime_type)
    if kind == "pdf":
        return _pdf(path)
    if kind == "text":
        return _text(path)
    if kind == "docx":
        return _docx(path)
    return _image(path)


def _pdf(path: Path) -> Extraction:
    try:
        import pymupdf  # type: ignore

        doc = pymupdf.open(str(path))
        pages = []
        for i, page in enumerate(doc, start=1):
            txt = (page.get_text("text") or "").strip()
            pages.append(Page(i, txt, len(txt), needs_ocr=len(txt) < 24))
        doc.close()
        return _finish(Extraction("pdf", pages, engine="pymupdf"))
    except ImportError:
        pass
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            txt = (page.extract_text() or "").strip()
            pages.append(Page(i, txt, len(txt), needs_ocr=len(txt) < 24))
        return _finish(Extraction("pdf", pages, engine="pypdf"))
    except Exception as exc:
        return _finish(Extraction("pdf", [], engine="none",
                                  warnings=[f"pdf text extraction failed: {type(exc).__name__}"]))


def _text(path: Path) -> Extraction:
    raw = path.read_text(encoding="utf-8", errors="replace")
    # Paginate long plain text so page references stay meaningful.
    per_page = 3000
    blocks = [raw[i:i + per_page] for i in range(0, max(len(raw), 1), per_page)] or [""]
    pages = [Page(i, b.strip(), len(b.strip())) for i, b in enumerate(blocks, start=1)]
    return _finish(Extraction("text", pages, engine="builtin"))


def _docx(path: Path) -> Extraction:
    try:
        import docx  # type: ignore
    except Exception as exc:
        return _finish(Extraction("docx", [], engine="none",
                                  warnings=[f"python-docx unavailable: {type(exc).__name__}"]))
    document = docx.Document(str(path))
    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    joined = "\n".join(parts)
    per_page = 3000
    blocks = [joined[i:i + per_page] for i in range(0, max(len(joined), 1), per_page)] or [""]
    pages = [Page(i, b.strip(), len(b.strip())) for i, b in enumerate(blocks, start=1)]
    return _finish(Extraction("docx", pages, engine="python-docx"))


def _image(path: Path) -> Extraction:
    """A picture has no text layer; OCR fills it in during ingest if available."""
    return _finish(Extraction("image", [Page(1, "", 0, needs_ocr=True)], engine="none",
                              warnings=["image document: text comes from OCR if a provider is ready"]))


def _finish(e: Extraction) -> Extraction:
    e.page_count = len(e.pages)
    e.total_chars = sum(p.char_count for p in e.pages)
    if e.page_count and all(p.needs_ocr for p in e.pages) and e.kind == "pdf":
        e.warnings.append("no text layer found — this looks like a scanned PDF; OCR is needed")
    return e
